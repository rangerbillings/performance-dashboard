"""
Live Web Dashboard - Performance Dashboard (Upload-Based Version)
Run with: streamlit run dashboard_app.py

Requires: pip install streamlit plotly pandas openpyxl
"""

import re
import io
import openpyxl
import pandas as pd
import streamlit as st
import plotly.graph_objects as go

# ------------------------------------------------------------------
# CONFIG
# ------------------------------------------------------------------
TARGETS = {"CC8D": 5.1, "FTC %": 75.0, "RPD": 780.0}
METRICS = ["Assigned Hours", "CC8D", "FTC %", "RPD", "Bodewell Engagement %"]

TARGET_METRIC_ALIASES = {
    "assigned hours": "Assigned Hours",
    "cc8d": "CC8D",
    "ftc %": "FTC %",
    "ftc%": "FTC %",
    "rpd": "RPD",
    "bodewell engagement %": "Bodewell Engagement %",
    "bodewell engagement%": "Bodewell Engagement %",
}

st.set_page_config(page_title="Performance Dashboard", layout="wide")


# ------------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------------
def extract_area_code(filename):
    match = re.search(r"\b(\d{3}-\d{2})\b", filename)
    return match.group(1) if match else None


# ------------------------------------------------------------------
# PARSER
# ------------------------------------------------------------------
@st.cache_data
def parse_weekly_workbook(file_bytes, filename):
    wb_in = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws_in = wb_in.worksheets[0]

    header_row, name_col, rzt_col = None, None, None
    for r in range(1, 8):
        for c in range(1, 25):
            val = ws_in.cell(row=r, column=c).value
            if isinstance(val, str):
                if val.strip().lower() == "tech name":
                    header_row, name_col = r, c
                if val.strip().upper() == "RZT":
                    rzt_col = c
        if header_row:
            break

    if header_row is None or name_col is None:
        raise ValueError("Could not locate a 'Tech Name' column header in the uploaded file.")

    metric_col = None
    for r in range(header_row + 1, min(header_row + 40, ws_in.max_row + 1)):
        for c in range(name_col + 1, name_col + 5):
            val = ws_in.cell(row=r, column=c).value
            if isinstance(val, str) and val.strip().lower() in TARGET_METRIC_ALIASES:
                metric_col = c
                break
        if metric_col:
            break
    if metric_col is None:
        metric_col = name_col + 1

    week_start_col = metric_col + 1
    week_labels = []
    c = week_start_col
    while True:
        val = ws_in.cell(row=header_row, column=c).value
        if val is None or str(val).strip() == "":
            break
        week_labels.append(str(val).strip())
        c += 1
    week_count = len(week_labels)

    if week_count == 0:
        raise ValueError("No week columns were detected in the header row.")

    tech_data = {}
    tech_id_map = {}
    current_tech = None

    for r in range(header_row + 1, ws_in.max_row + 1):
        name_val = ws_in.cell(row=r, column=name_col).value
        if name_val is not None and str(name_val).strip() != "":
            current_tech = str(name_val).strip()
            if rzt_col:
                rzt_val = ws_in.cell(row=r, column=rzt_col).value
                tech_id_map[current_tech] = str(rzt_val).strip() if rzt_val is not None else ""
            if current_tech not in tech_data:
                tech_data[current_tech] = {m: [None] * week_count for m in METRICS}

        if current_tech is None:
            continue

        metric_val = ws_in.cell(row=r, column=metric_col).value
        if metric_val is None or str(metric_val).strip() == "":
            continue

        standard_metric = TARGET_METRIC_ALIASES.get(str(metric_val).strip().lower())
        if standard_metric is None:
            continue

        row_values = []
        for i in range(week_count):
            cell = ws_in.cell(row=r, column=week_start_col + i)
            v = cell.value
            if v is None:
                row_values.append(None)
                continue
            if isinstance(v, (int, float)):
                fmt = cell.number_format or ""
                row_values.append(round(v * 100, 2) if "%" in fmt else round(float(v), 2))
            else:
                text_val = str(v).strip().replace("$", "").replace(",", "")
                if text_val.endswith("%"):
                    try:
                        row_values.append(round(float(text_val[:-1]), 2))
                    except ValueError:
                        row_values.append(None)
                else:
                    try:
                        row_values.append(round(float(text_val), 2))
                    except ValueError:
                        row_values.append(None)

        tech_data[current_tech][standard_metric] = row_values

    return tech_data, week_labels, tech_id_map


# ------------------------------------------------------------------
# CALCULATION HELPERS
# ------------------------------------------------------------------
def safe_avg(values):
    clean = [v for v in values if v is not None]
    return round(sum(clean) / len(clean), 2) if clean else None


def week_team_average(tech_data, techs, metric, week_idx):
    values = [tech_data[t][metric][week_idx] for t in techs if tech_data[t][metric][week_idx] is not None]
    return round(sum(values) / len(values), 2) if values else None


def tech_overall_average(tech_data, tech, metric):
    return safe_avg(tech_data[tech][metric])


def full_rank(tech_data, techs, metric):
    averaged = [(t, tech_overall_average(tech_data, t, metric)) for t in techs]
    averaged = [(t, v) for t, v in averaged if v is not None]
    averaged.sort(key=lambda x: x[1], reverse=True)
    return {t: i + 1 for i, (t, v) in enumerate(averaged)}


def rank_techs(tech_data, techs, metric, top_n=3):
    averaged = [(t, tech_overall_average(tech_data, t, metric)) for t in techs]
    averaged = [(t, v) for t, v in averaged if v is not None]
    averaged.sort(key=lambda x: x[1], reverse=True)
    return averaged[:top_n], averaged[-top_n:][::-1]


def fmt_val(metric, value):
    if value is None:
        return "N/A"
    if metric == "CC8D":
        return f"{value:.2f}"
    if metric in ("FTC %", "Bodewell Engagement %"):
        return f"{value:.1f}%"
    if metric == "RPD":
        return f"${value:,.2f}"
    if metric == "Assigned Hours":
        return f"{value:.1f}"
    return str(value)


def make_trend_chart(weeks, series_dict, metric, target=None):
    fig = go.Figure()
    for name, values in series_dict.items():
        fig.add_trace(go.Scatter(
            x=weeks, y=values, mode="lines+markers+text",
            name=name, text=[fmt_val(metric, v) if v is not None else "" for v in values],
            textposition="top center"
        ))
    if target is not None:
        fig.add_trace(go.Scatter(
            x=weeks, y=[target] * len(weeks), mode="lines",
            name=f"Target ({target})", line=dict(dash="dash", color="red")
        ))
    fig.update_layout(
        title=f"{metric} - Weekly Trend", xaxis_title="Week", yaxis_title=metric,
        height=400, margin=dict(t=50, b=30)
    )
    return fig


# ------------------------------------------------------------------
# APP START
# ------------------------------------------------------------------
st.title("📊 Performance Dashboard")
st.caption("Upload your weekly Tech Stats export(s) below to get started.")

col_a, col_b = st.columns(2)
with col_a:
    upload_133_23 = st.file_uploader("Upload 133-23 Tech Stats file", type=["xlsx"], key="upload_23")
with col_b:
    upload_133_24 = st.file_uploader("Upload 133-24 Tech Stats file", type=["xlsx"], key="upload_24")

uploaded_files = {}
if upload_133_23 is not None:
    uploaded_files[upload_133_23.name] = upload_133_23
if upload_133_24 is not None:
    uploaded_files[upload_133_24.name] = upload_133_24

if not uploaded_files:
    st.info("👆 Upload at least one weekly Tech Stats Excel file above to view the dashboard.")
    st.stop()

file_labels = list(uploaded_files.keys())
selected_label = st.selectbox("Select data source to view:", file_labels)
selected_file = uploaded_files[selected_label]
area_code = extract_area_code(selected_label) or ""

file_bytes = selected_file.getvalue()

try:
    tech_data, WEEKS, tech_id_map = parse_weekly_workbook(file_bytes, selected_label)
except Exception as e:
    st.error(f"Could not parse this file: {e}")
    st.stop()

TECHS = list(tech_data.keys())
st.caption(f"Area: **{area_code or 'N/A'}**  |  {len(TECHS)} Technicians  |  "
           f"Period: {WEEKS[0]} to {WEEKS[-1]}  |  Source: {selected_label}")

tab1, tab2, tab3, tab4, tab5 = st.tabs(
    ["📋 Executive Summary", "🔍 Deep Dive", "⚖️ Comparison", "🏆 Rankings", "📈 Trends"]
)

# ------------------------------------------------------------------
# TAB 1: EXECUTIVE SUMMARY
# ------------------------------------------------------------------
with tab1:
    st.subheader("Top-Line Results")
    col1, col2, col3 = st.columns(3)

    cc8d_avg = safe_avg([week_team_average(tech_data, TECHS, "CC8D", i) for i in range(len(WEEKS))])
    ftc_avg = safe_avg([week_team_average(tech_data, TECHS, "FTC %", i) for i in range(len(WEEKS))])
    rpd_avg = safe_avg([week_team_average(tech_data, TECHS, "RPD", i) for i in range(len(WEEKS))])

    with col1:
        delta = round(cc8d_avg - TARGETS["CC8D"], 2) if cc8d_avg else None
        st.metric("Team Avg CC8D", fmt_val("CC8D", cc8d_avg),
                   delta=f"{delta:+.2f} vs target" if delta is not None else None)
    with col2:
        delta = round(ftc_avg - TARGETS["FTC %"], 1) if ftc_avg else None
        st.metric("Team Avg FTC%", fmt_val("FTC %", ftc_avg),
                   delta=f"{delta:+.1f} pts vs target" if delta is not None else None)
    with col3:
        delta = round(rpd_avg - TARGETS["RPD"], 2) if rpd_avg else None
        st.metric("Team Avg RPD", fmt_val("RPD", rpd_avg),
                   delta=f"{delta:+,.2f} vs target" if delta is not None else None)

    st.subheader("Target Attainment")
    for metric in ["CC8D", "FTC %", "RPD"]:
        target = TARGETS[metric]
        meeting = [(t, tech_overall_average(tech_data, t, metric)) for t in TECHS]
        meeting = [(t, v) for t, v in meeting if v is not None]
        above = sorted([x for x in meeting if x[1] >= target], key=lambda x: -x[1])
        below = sorted([x for x in meeting if x[1] < target], key=lambda x: x[1])

        st.markdown(f"**{metric} (Target: {target})**")
        c1, c2 = st.columns(2)
        with c1:
            st.success("Meeting/Exceeding: " + (", ".join(f"{t} ({fmt_val(metric, v)})" for t, v in above) or "None"))
        with c2:
            st.error("Below Target: " + (", ".join(f"{t} ({fmt_val(metric, v)})" for t, v in below) or "None"))

    st.subheader("RPD - Top 3 / Bottom 3 Performers")
    top_rpd, bottom_rpd = rank_techs(tech_data, TECHS, "RPD", top_n=3)
    c1, c2 = st.columns(2)
    with c1:
        st.success("**Top 3:** " + ", ".join(f"{t} ({fmt_val('RPD', v)})" for t, v in top_rpd))
    with c2:
        st.error("**Bottom 3:** " + ", ".join(f"{t} ({fmt_val('RPD', v)})" for t, v in bottom_rpd))

    st.subheader("Top Performers (Multi-Metric)")
    counts = {}
    for metric in ["CC8D", "FTC %", "RPD", "Bodewell Engagement %"]:
        top, _ = rank_techs(tech_data, TECHS, metric, top_n=3)
        for t, _ in top:
            counts[t] = counts.get(t, 0) + 1
    leaders = sorted([(t, c) for t, c in counts.items() if c >= 2], key=lambda x: -x[1])
    if leaders:
        st.success(", ".join(f"{t} (top 3 in {c}/4 metrics)" for t, c in leaders))
    else:
        st.info("No standout multi-metric leaders this period.")

    st.subheader("Priority Concerns (Multi-Metric)")
    counts = {}
    for metric in ["CC8D", "FTC %", "RPD"]:
        _, bottom = rank_techs(tech_data, TECHS, metric, top_n=3)
        for t, _ in bottom:
            counts[t] = counts.get(t, 0) + 1
    laggards = sorted([(t, c) for t, c in counts.items() if c >= 2], key=lambda x: -x[1])
    if laggards:
        for t, c in laggards:
            st.error(f"{t} - Bottom 3 in {c}/3 key metrics")
    else:
        st.info("No multi-metric concerns flagged this period.")


# ------------------------------------------------------------------
# TAB 2: DEEP DIVE
# ------------------------------------------------------------------
with tab2:
    selected_tech = st.selectbox("Select Technician:", TECHS, key="deepdive_tech")
    tech_id = tech_id_map.get(selected_tech, "N/A")
    st.caption(f"Tech ID (RZT): {tech_id}")

    rank_maps = {m: full_rank(tech_data, TECHS, m) for m in METRICS}

    cols = st.columns(len(METRICS))
    for i, metric in enumerate(METRICS):
        avg = tech_overall_average(tech_data, selected_tech, metric)
        rank = rank_maps[metric].get(selected_tech, "N/A")
        with cols[i]:
            st.metric(metric, fmt_val(metric, avg), delta=f"Rank {rank}/{len(TECHS)}")

    st.divider()
    for metric in METRICS:
        target = TARGETS.get(metric)
        fig = make_trend_chart(WEEKS, {selected_tech: tech_data[selected_tech][metric]}, metric, target=target)
        st.plotly_chart(fig, use_container_width=True)


# ------------------------------------------------------------------
# TAB 3: COMPARISON
# ------------------------------------------------------------------
with tab3:
    c1, c2 = st.columns(2)
    with c1:
        tech_a = st.selectbox("Technician A:", TECHS, index=0, key="comp_a")
    with c2:
        tech_b = st.selectbox("Technician B:", TECHS, index=min(1, len(TECHS) - 1), key="comp_b")

    for metric in METRICS:
        target = TARGETS.get(metric)
        series = {tech_a: tech_data[tech_a][metric], tech_b: tech_data[tech_b][metric]}
        fig = make_trend_chart(WEEKS, series, metric, target=target)
        st.plotly_chart(fig, use_container_width=True)

        avg_a = tech_overall_average(tech_data, tech_a, metric)
        avg_b = tech_overall_average(tech_data, tech_b, metric)
        if avg_a is not None and avg_b is not None:
            leader = tech_a if avg_a >= avg_b else tech_b
            gap = abs(avg_a - avg_b)
            st.caption(f"**{leader}** leads by {fmt_val(metric, gap)} on average ({metric})")


# ------------------------------------------------------------------
# TAB 4: RANKINGS
# ------------------------------------------------------------------
with tab4:
    for metric in METRICS:
        st.subheader(metric)
        top, bottom = rank_techs(tech_data, TECHS, metric, top_n=5)
        col1, col2 = st.columns(2)
        with col1:
            st.markdown("**Top 5**")
            df_top = pd.DataFrame(top, columns=["Technician", "Average"])
            df_top["Average"] = df_top["Average"].apply(lambda v: fmt_val(metric, v))
            st.dataframe(df_top, hide_index=True, use_container_width=True)
        with col2:
            st.markdown("**Bottom 5**")
            df_bottom = pd.DataFrame(bottom, columns=["Technician", "Average"])
            df_bottom["Average"] = df_bottom["Average"].apply(lambda v: fmt_val(metric, v))
            st.dataframe(df_bottom, hide_index=True, use_container_width=True)


# ------------------------------------------------------------------
# TAB 5: TRENDS
# ------------------------------------------------------------------
with tab5:
    for metric in METRICS:
        target = TARGETS.get(metric)
        team_avg = [week_team_average(tech_data, TECHS, metric, i) for i in range(len(WEEKS))]
        fig = make_trend_chart(WEEKS, {"Team Average": team_avg}, metric, target=target)
        st.plotly_chart(fig, use_container_width=True)
